import pandas as pd
import glob
from datetime import datetime
from openpyxl.styles import Font,PatternFill,Border,Side,Alignment
#หาไฟล์ทั้งหมดในโฟลเดอร์
files = glob.glob('input/*.xlsx')
print(f"พบ {len(files)} ไฟล์")
#อ่านข้อมูลแล้วดึงข้อมูลไปเก็บในลิสต์
all_data = []
for file in files:
    df = pd.read_excel(file)
    all_data.append(df)
#รวมไฟล์ทุกไฟล์เป็นตชตารางเดียว
combined_data = pd.concat(all_data, ignore_index=True)
#แปลงคอลัมน์ dete ให้เป็นวันที่จริงๆ
combined_data['date'] = pd.to_datetime(combined_data['date']).dt.date
#สรุปยอดขายแยกตามสาขา
total_sales = combined_data.groupby('branch')['total'].sum().reset_index()
total_sales.columns = ["สาขา", "total"]
#หาTop 20 สินค้าที่ขายดีที่สุด
top20 = combined_data.groupby(["product_code", "product_name"]).agg(
    total_qty=("qty", "sum"),
    total_amount=("total", "sum")
).reset_index().sort_values("total_qty", ascending=False).head(20)

#ตั้งชื่อไฟล์ตามวันปัจจุบัน
file_name = f'output/Daily_Report_{datetime.today().strftime("%Y-%m-%d")}.xlsx'
#บันทึกลง Excel 
with pd.ExcelWriter(file_name) as writer:
    #กำหนด style เส้น
    thin= Side(style='thin')
    medium= Side(style='medium')
    #สร้างชีตใหม่
    combined_data.to_excel(writer,sheet_name='all_สาขา',index=False)
    worksheet = writer.sheets['all_สาขา']
    top20.to_excel(writer,sheet_name='Top20', index=False)
    total_sales.to_excel(writer,sheet_name="Summary_Branch",index=False)
    #ตกแต่งหัวตาราง
    for sheet_name in ['all_สาขา','Top20','Summary_Branch']:
        ws = writer.sheets[sheet_name]

        for col in ['A','B','C','D','E','F','G','H']:
            ws.column_dimensions[col].width = 15
        #ล็อคหัวตาราง
        ws.freeze_panes='A2'
        #ตกแต่งข้อความและเส้นหัวตาราง
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill= PatternFill("solid",start_color="00B050")
    
            cell.border = Border(
                left=medium,right=medium,
                top=medium,bottom=medium
            )
            #ทำให้หัวตารางอยู่กึ่งกลาง
            cell.alignment = Alignment(horizontal="center")
        #ตกแต่ง
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                #เพิ่ม ,
                if isinstance(cell.value,(int,float)):
                    cell.number_format = '#,##0'
                    #จัดตัวเลขให้อยู่ทางขวา
                    cell.alignment = Alignment(horizontal="right")
                else:
                    #จัดข้อมูลอื่นๆให้อยู่ทางซ้าย
                    cell.alignment = Alignment(horizontal="left")
                #ตกแต่งเส้นตาราง
                
                cell.border = Border(
                left=thin,right=thin,
                top=thin,bottom=thin
            )      
